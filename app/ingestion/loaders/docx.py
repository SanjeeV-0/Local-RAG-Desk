from pathlib import Path

from docx import Document as DocxDocument
from docx.oxml.table import CT_Tbl
from docx.oxml.text.paragraph import CT_P
from docx.table import Table as DocxTable
from docx.text.paragraph import Paragraph as DocxParagraph

from app.ingestion.loaders.base import DocumentLoader, LoadedPage
from app.ingestion.structure import TableData
from app.ingestion.table_serializer import table_to_markdown


class DocxLoader(DocumentLoader):
    """Load Microsoft Word (.docx) documents as a single logical page.

    Extracts body content in TRUE DOCUMENT ORDER -- paragraphs and tables
    interleaved exactly as they appear in the file, by walking
    `document.element.body`'s raw XML children directly (`CT_P`/`CT_Tbl`)
    rather than python-docx's `document.paragraphs`/`document.tables`,
    which are two separate flat sequences that each discard the other
    element type's position. That separation previously meant every table
    was rendered after ALL paragraph text regardless of where it actually
    appeared, which in turn meant `StructureExtractor` (which tracks the
    current heading by walking the text top-to-bottom) attributed every
    table to whichever heading happened to be LAST in the whole document --
    silently wrong for any document with more than one heading-with-table.
    Walking the body in order fixes this: a table between "## Results" and
    "## Conclusion" now stays between them in the emitted content, so it is
    correctly attributed to "Results".

    Word documents don't carry reliable page-break metadata without
    rendering, so (like MarkdownLoader) no attempt is made to infer
    physical page numbers -- the whole document is returned as one
    LoadedPage.

    Each table's genuinely structured data (python-docx already exposes
    real rows/cells, unlike PDF text extraction) is rendered as a standard
    Markdown pipe table -- via the same `table_serializer.table_to_markdown`
    used elsewhere in ingestion -- as its own block, separated from
    surrounding paragraph text by a blank line so `StructureExtractor`
    always sees a clean boundary. Rendering as real pipe-table syntax
    (header row + `|---|---|` separator) is what lets `StructureExtractor`
    detect these as genuine TABLE structural units downstream -- the first
    row of each table is treated as its header row.

    Body-level XML elements that are neither a paragraph nor a table (e.g.
    `CT_SectPr`, section properties) are silently skipped -- same as
    before, when only `document.paragraphs`/`document.tables` were read and
    nothing else was ever consulted.
    """

    def load(self, path: Path) -> list[LoadedPage]:
        document = DocxDocument(path)

        blocks: list[str] = []
        current_paragraph_lines: list[str] = []
        table_index = 0

        def flush_paragraphs() -> None:
            if current_paragraph_lines:
                blocks.append("\n".join(current_paragraph_lines))
                current_paragraph_lines.clear()

        for child in document.element.body.iterchildren():
            if isinstance(child, CT_P):
                paragraph = DocxParagraph(child, document)
                text = paragraph.text.strip()

                if not text:
                    continue

                style_name = paragraph.style.name if paragraph.style is not None else ""

                if style_name.startswith("Heading "):
                    try:
                        level = int(style_name.removeprefix("Heading ").strip())
                    except ValueError:
                        level = 0

                    if 1 <= level <= 6:
                        text = f"{'#' * level} {text}"

                current_paragraph_lines.append(text)

            elif isinstance(child, CT_Tbl):
                # A table always starts its own block -- flush whatever
                # paragraph run preceded it first, so the two never merge
                # into one run-on block with no boundary between them.
                flush_paragraphs()

                table_index += 1

                table = DocxTable(child, document)
                rows = [[cell.text.strip() for cell in row.cells] for row in table.rows]
                rows = [row for row in rows if any(cell for cell in row)]

                if not rows:
                    continue

                headers, *data_rows = rows

                table_data = TableData(
                    table_id=f"docx-table-{table_index}",
                    # .docx has no reliable "this is the title" signal
                    # separate from a preceding paragraph -- never
                    # fabricate one.
                    title=None,
                    headers=headers,
                    rows=data_rows,
                    page_numbers=[1],
                )

                blocks.append(table_to_markdown(table_data))

            # Any other body-level element (e.g. CT_SectPr) is silently
            # skipped -- see class docstring.

        flush_paragraphs()

        # Blank-line-separated so StructureExtractor always sees a clean
        # boundary between any two adjacent blocks, regardless of what kind
        # either one is.
        content = "\n\n".join(block for block in blocks if block)

        return [LoadedPage(page_number=1, content=content)]
